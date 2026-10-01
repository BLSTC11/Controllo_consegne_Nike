"""Controllo consegne stagionali Nike.

Confronta l'ordine confermato (ordini_confermati/*.xlsx) con TUTTI i file di consegna
(consegne/*.xlsx), abbinando per EAN + Stagione, e produce in output/:
  - controllo_consegne_<stagione>.xlsx  (riepilogo, eccedenze, non ordinato, residui, dettaglio)
  - controllo_consegne_<stagione>.pdf   (report stampabile, generato con Google Chrome)

Uso: python3 controllo_consegne.py
"""
from datetime import date
from html import escape
from pathlib import Path
import re
import subprocess

import pandas as pd

BASE = Path(__file__).parent
ORD_DIR, CONS_DIR, OUT_DIR = BASE / "ordini_confermati", BASE / "consegne", BASE / "output"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# Codici stagione Nike nei file di consegna -> nome stagione dell'ordine.
# Se un codice non è qui, viene ricavato dal suffisso del Truck (es. "BGRO093/FA26").
SEASON_CODES = {263: "FA26"}


def xlsx(folder):
    return sorted(p for p in folder.glob("*.xlsx") if not p.name.startswith("~$"))


# ---------------------------------------------------------------- caricamento
def load_order():
    frames = []
    for f in xlsx(ORD_DIR):
        df = pd.read_excel(f, header=1)
        df = df[df["EAN UPC Cd"].notna()].copy()
        df["File ordine"] = f.name
        frames.append(df)
    o = pd.concat(frames, ignore_index=True)
    o = o.rename(columns={"EAN UPC Cd": "EAN", "Qty status NNT": "Qty ordinata",
                          "FPC Price IN EUR w/o VAT": "Prezzo FPC", "PO name SAP": "PO",
                          "Article Number": "Articolo", "Article Name": "Descrizione",
                          "Color Desc": "Colore", "Size US": "Taglia US", "EU Size": "Taglia EU",
                          "Season": "Stagione"})
    o["EAN"] = o["EAN"].astype("int64")
    o["PO"] = o["PO"].str.upper()
    o["Stagione"] = o["Stagione"].astype(str).str.strip().str.upper()
    o["Taglia US"] = o["Taglia US"].astype(str)
    o["Valore ordinato"] = o["Qty ordinata"] * o["Prezzo FPC"]
    keys = ["EAN", "Stagione", "PO", "BU", "Articolo", "Descrizione", "Colore", "Taglia US",
            "Taglia EU", "Gender", "Prezzo FPC"]
    return o.groupby(keys, as_index=False, dropna=False)[["Qty ordinata", "Valore ordinato"]].sum()


def delivery_season(row):
    code = row["Season"]
    if pd.notna(code) and int(code) in SEASON_CODES:
        return SEASON_CODES[int(code)]
    m = re.search(r"/([A-Z]{2}\d{2})\s*$", str(row.get("Truck", "")))
    return m.group(1) if m else f"COD {code}"


def load_deliveries():
    frames = []
    for f in xlsx(CONS_DIR):
        d = pd.read_excel(f)
        d = d[d["Barcode"].notna()].copy()
        d["File consegna"] = f.name
        frames.append(d)
    if not frames:
        raise SystemExit("Nessun file di consegna in consegne/")
    d = pd.concat(frames, ignore_index=True)
    d["EAN"] = d["Barcode"].astype("int64")
    d["Stagione"] = d.apply(delivery_season, axis=1)
    d["Valore FPC"] = d["Dlv.qty"] * d["FPC Price w/o VAT in EUR"]
    d["Po name"] = d["Po name"].str.upper()
    return d


# ---------------------------------------------------------------- confronto
def compare(o, d):
    per = d.groupby(["EAN", "Stagione"], as_index=False).agg(
        **{"Qty consegnata": ("Dlv.qty", "sum"),
           "Valore FPC consegnato": ("Valore FPC", "sum"),
           "Valore netto consegnato": ("Net value", "sum"),
           "Prezzo FPC consegna": ("FPC Price w/o VAT in EUR", "max"),
           "PO consegna": ("Po name", lambda s: ", ".join(sorted(set(s)))),
           "File consegna": ("File consegna", lambda s: ", ".join(sorted(set(s)))),
           "Articolo cons.": ("Material", "first"),
           "Descrizione cons.": ("Material description", "first"),
           "Colore cons.": ("Color", "first"),
           "Taglia cons.": ("Size", "first"),
           "BU cons.": ("Division", "first")})

    m = o.merge(per, on=["EAN", "Stagione"], how="outer", indicator=True)
    for c in ["Qty ordinata", "Valore ordinato", "Qty consegnata", "Valore FPC consegnato",
              "Valore netto consegnato"]:
        m[c] = m[c].fillna(0)

    # Righe consegnate senza corrispondenza EAN+Stagione: completa i dati descrittivi dalla consegna
    only_d = m["_merge"] == "right_only"
    for a, b in [("Articolo", "Articolo cons."), ("Descrizione", "Descrizione cons."),
                 ("Colore", "Colore cons."), ("Taglia US", "Taglia cons."), ("BU", "BU cons.")]:
        m.loc[only_d, a] = m.loc[only_d, b]
    m["Taglia US"] = m["Taglia US"].astype(str)
    ord_ean_season = o.groupby("EAN")["Stagione"].agg(lambda s: ", ".join(sorted(set(s))))
    m["Stagione ordine EAN"] = m["EAN"].map(ord_ean_season)

    m["Qty consegnata su ordine"] = m[["Qty ordinata", "Qty consegnata"]].min(axis=1)
    m["Qty residua"] = (m["Qty ordinata"] - m["Qty consegnata"]).clip(lower=0)
    m["Qty in eccesso"] = (m["Qty consegnata"] - m["Qty ordinata"]).clip(lower=0).where(~only_d, 0)
    pr = m["Prezzo FPC"].fillna(m["Prezzo FPC consegna"])
    m["Valore consegnato su ordine"] = m["Qty consegnata su ordine"] * pr
    m["Valore residuo"] = m["Qty residua"] * pr
    m["Valore eccesso"] = m["Qty in eccesso"] * pr

    def stato(r):
        if r["_merge"] == "right_only":
            return "STAGIONE DIVERSA" if pd.notna(r["Stagione ordine EAN"]) else "NON ORDINATO"
        if r["Qty consegnata"] == 0:
            return "Da consegnare"
        if r["Qty in eccesso"] > 0:
            return "ECCEDENZA"
        return "Parziale" if r["Qty residua"] > 0 else "Completo"

    m["Stato"] = m.apply(stato, axis=1)
    both = m["_merge"] == "both"
    m["Diff. prezzo FPC"] = (m["Prezzo FPC consegna"] - m["Prezzo FPC"]).where(both & (m["Qty consegnata"] > 0))
    m["PO diverso"] = both & m["PO consegna"].notna() & (m["PO consegna"] != m["PO"])
    return m.drop(columns=["_merge", "Articolo cons.", "Descrizione cons.", "Colore cons.",
                           "Taglia cons.", "BU cons."])


def summaries(m, d):
    ok = m[~m["Stato"].isin(["NON ORDINATO", "STAGIONE DIVERSA"])]
    extra = m[m["Stato"].isin(["NON ORDINATO", "STAGIONE DIVERSA"])]
    v_ord = ok["Valore ordinato"].sum()
    tot = pd.DataFrame([
        ["Ordinato (confermato)", ok["Qty ordinata"].sum(), v_ord],
        ["Consegnato sull'ordine", ok["Qty consegnata su ordine"].sum(), ok["Valore consegnato su ordine"].sum()],
        ["Residuo da consegnare", ok["Qty residua"].sum(), ok["Valore residuo"].sum()],
        ["Consegnato in eccedenza", ok["Qty in eccesso"].sum(), ok["Valore eccesso"].sum()],
        ["Consegnato NON ordinato", extra.loc[extra.Stato == "NON ORDINATO", "Qty consegnata"].sum(),
         extra.loc[extra.Stato == "NON ORDINATO", "Valore FPC consegnato"].sum()],
        ["Consegnato di altra stagione", extra.loc[extra.Stato == "STAGIONE DIVERSA", "Qty consegnata"].sum(),
         extra.loc[extra.Stato == "STAGIONE DIVERSA", "Valore FPC consegnato"].sum()],
        ["Totale consegnato (tutti i file)", d["Dlv.qty"].sum(), d["Valore FPC"].sum()],
    ], columns=["Voce", "Pezzi", "Valore FPC €"])
    tot["% su ordinato (valore)"] = tot["Valore FPC €"] / v_ord

    def group(by):
        g = ok.groupby(by).agg(**{"Pezzi ordinati": ("Qty ordinata", "sum"),
                                  "Pezzi consegnati": ("Qty consegnata su ordine", "sum"),
                                  "Pezzi residui": ("Qty residua", "sum"),
                                  "Pezzi eccedenza": ("Qty in eccesso", "sum"),
                                  "Valore ordinato €": ("Valore ordinato", "sum"),
                                  "Valore consegnato €": ("Valore consegnato su ordine", "sum"),
                                  "Valore residuo €": ("Valore residuo", "sum")})
        g["% consegnato (valore)"] = g["Valore consegnato €"] / g["Valore ordinato €"]
        return g.reset_index()

    per_file = d.groupby(["File consegna", "Stagione", "Packing List", "Truck", "Po name"], as_index=False).agg(
        **{"Pezzi": ("Dlv.qty", "sum"), "Valore FPC €": ("Valore FPC", "sum"),
           "Valore netto €": ("Net value", "sum")})
    per_file["Packing List"] = per_file["Packing List"].astype("int64").astype(str)

    art = ok.groupby(["Stagione", "BU", "Articolo", "Descrizione", "Colore", "Prezzo FPC"], as_index=False).agg(
        **{"Pezzi ordinati": ("Qty ordinata", "sum"), "Pezzi consegnati": ("Qty consegnata su ordine", "sum"),
           "Pezzi residui": ("Qty residua", "sum"), "Pezzi eccedenza": ("Qty in eccesso", "sum"),
           "Valore ordinato €": ("Valore ordinato", "sum"), "Valore residuo €": ("Valore residuo", "sum")})
    art["% consegnato"] = 1 - art["Valore residuo €"] / art["Valore ordinato €"]
    return tot, group(["Stagione", "BU"]), group(["Stagione", "PO"]), per_file, art


# ---------------------------------------------------------------- output
DET_COLS = ["Stato", "Stagione", "Stagione ordine EAN", "PO", "BU", "Articolo", "Descrizione", "Colore",
            "Taglia US", "Taglia EU", "EAN", "Prezzo FPC", "Qty ordinata", "Qty consegnata", "Qty residua",
            "Qty in eccesso", "Valore ordinato", "Valore consegnato su ordine", "Valore residuo",
            "Valore eccesso", "Valore FPC consegnato", "Valore netto consegnato", "Prezzo FPC consegna",
            "Diff. prezzo FPC", "PO consegna", "PO diverso", "File consegna"]


def write_excel(path, tot, bu, po, per_file, art, m):
    from openpyxl.styles import Font, PatternFill
    det = m[DET_COLS].sort_values(["Stagione", "BU", "Articolo", "EAN"])
    sheets = [
        ("Consegne", per_file),
        ("Eccedenze", det[det.Stato == "ECCEDENZA"]),
        ("Non ordinato", det[det.Stato.isin(["NON ORDINATO", "STAGIONE DIVERSA"])]),
        ("Differenze prezzo-PO", det[(det["Diff. prezzo FPC"].fillna(0).abs() > 0.009) | det["PO diverso"]]),
        ("Residuo per articolo", art[art["Pezzi residui"] > 0].sort_values("Valore residuo €", ascending=False)),
        ("Per articolo", art.sort_values(["BU", "Articolo"])),
        ("Dettaglio EAN", det),
    ]
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDE4EE")
    header_rows = {}  # foglio -> righe d'intestazione (1-based)
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        r, hdrs = 0, []
        for title, df in [("Totali stagione", tot), ("Per reparto", bu), ("Per numero d'ordine", po)]:
            df.to_excel(w, sheet_name="Riepilogo", startrow=r + 1, index=False)
            w.book["Riepilogo"].cell(row=r + 1, column=1, value=title).font = Font(bold=True, size=12)
            hdrs.append(r + 2)
            r += len(df) + 4
        header_rows["Riepilogo"] = hdrs
        for name, df in sheets:
            df.to_excel(w, sheet_name=name, index=False)
            header_rows[name] = [1]
        for ws in w.book.worksheets:
            hr = header_rows[ws.title]
            if ws.title != "Riepilogo":
                ws.freeze_panes = "A2"
            pct_cols = set()
            for h in hr:
                for c in ws[h]:
                    if c.value is not None:
                        c.font, c.fill = bold, fill
                        if "%" in str(c.value):
                            pct_cols.add((h, c.column))
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, float):
                        h = max(x for x in hr if x < c.row) if any(x < c.row for x in hr) else None
                        c.number_format = "0.0%" if (h, c.column) in pct_cols else "#,##0.00"
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(
                    max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2, 40)


def fmt(v, col=""):
    if isinstance(v, float) and pd.isna(v):
        return ""
    if "%" in col:
        return f"{v:.1%}".replace(".", ",")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if float(v).is_integer() and "€" not in col and "Valore" not in col and "Prezzo" not in col:
            return f"{int(v):,}".replace(",", ".")
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return escape(str(v))


def table(df, cols=None, max_rows=None):
    cols = cols or list(df.columns)
    rows = df[cols] if max_rows is None else df[cols].head(max_rows)
    h = "".join(f"<th>{escape(c)}</th>" for c in cols)
    b = "".join("<tr>" + "".join(
        f"<td class='{'n' if isinstance(v, (int, float)) and not isinstance(v, bool) else ''}'>{fmt(v, c)}</td>"
        for c, v in zip(cols, r)) + "</tr>" for r in rows.itertuples(index=False))
    more = f"<p class='note'>… altre {len(df) - max_rows} righe nel file Excel.</p>" if max_rows and len(df) > max_rows else ""
    return f"<table><thead><tr>{h}</tr></thead><tbody>{b}</tbody></table>{more}"


def write_pdf(html_path, pdf_path, season, tot, bu, po, per_file, art, m):
    det_cols = ["Stagione", "BU", "Articolo", "Descrizione", "Taglia US", "EAN", "Qty ordinata",
                "Qty consegnata", "Qty in eccesso", "Valore eccesso"]
    exc = m[m.Stato == "ECCEDENZA"]
    non = m[m.Stato.isin(["NON ORDINATO", "STAGIONE DIVERSA"])]
    diff = m[(m["Diff. prezzo FPC"].fillna(0).abs() > 0.009) | m["PO diverso"]]
    res = art[art["Pezzi residui"] > 0].sort_values("Valore residuo €", ascending=False)

    def section(title, df, cols, empty, max_rows=None):
        body = table(df, cols, max_rows) if len(df) else f"<p class='ok'>✓ {empty}</p>"
        return f"<h2>{title}</h2>{body}"

    html = f"""<!doctype html><html><head><meta charset="utf-8"><title>Controllo consegne {season}</title>
<style>
@page {{ size: A4 landscape; margin: 12mm; }}
body {{ font-family: -apple-system, Helvetica, Arial, sans-serif; font-size: 9pt; color: #222; }}
h1 {{ font-size: 16pt; margin: 0 0 2mm; }} h2 {{ font-size: 12pt; margin: 6mm 0 2mm; border-bottom: 1px solid #999; }}
.sub {{ color: #555; margin-bottom: 4mm; }}
table {{ border-collapse: collapse; width: 100%; margin-bottom: 2mm; }}
th, td {{ border: 1px solid #ccc; padding: 1.2mm 1.6mm; text-align: left; }}
th {{ background: #dde4ee; }} td.n {{ text-align: right; white-space: nowrap; }}
tr {{ page-break-inside: avoid; }} thead {{ display: table-header-group; }}
.ok {{ color: #1a7f37; font-weight: bold; }} .note {{ color: #666; font-style: italic; }}
</style></head><body>
<h1>Controllo consegne Nike – stagione {escape(season)}</h1>
<div class="sub">Ordine: {escape(', '.join(sorted(set(xlsx_names(ORD_DIR)))))}<br>
Consegne analizzate: {len(per_file['File consegna'].unique())} file – {escape(', '.join(sorted(per_file['File consegna'].unique())))}<br>
Generato il {date.today():%d/%m/%Y}. Valori a prezzo FPC netto IVA; abbinamento per EAN + Stagione.</div>
<h2>Totali</h2>{table(tot)}
<h2>Per reparto</h2>{table(bu)}
<h2>Per numero d'ordine</h2>{table(po)}
<h2>Consegne ricevute</h2>{table(per_file)}
{section("Consegnato in eccedenza rispetto all'ordine", exc, det_cols, "Nessuna eccedenza.")}
{section("Consegnato ma non ordinato (o di altra stagione)", non, ["Stato", "Stagione", "Stagione ordine EAN", "BU", "Articolo", "Descrizione", "Taglia US", "EAN", "Qty consegnata", "Valore FPC consegnato", "File consegna"], "Nessun articolo consegnato fuori ordine.")}
{section("Differenze di prezzo o numero d'ordine", diff, ["Stagione", "Articolo", "Descrizione", "Taglia US", "EAN", "Prezzo FPC", "Prezzo FPC consegna", "PO", "PO consegna"], "Nessuna differenza di prezzo o PO.")}
{section("Residuo da consegnare per articolo", res, ["BU", "Articolo", "Descrizione", "Colore", "Pezzi ordinati", "Pezzi consegnati", "Pezzi residui", "Valore residuo €", "% consegnato"], "Ordine completamente consegnato.")}
</body></html>"""
    html_path.write_text(html, encoding="utf-8")
    try:
        subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf_path}", html_path.as_uri()],
                       check=True, capture_output=True, timeout=120)
        return True
    except Exception as e:  # Chrome assente: resta l'HTML, stampabile dal browser
        print("PDF non generato:", e)
        return False


def xlsx_names(folder):
    return [p.name for p in xlsx(folder)]


def main():
    o = load_order()
    d = load_deliveries()
    m = compare(o, d)
    tot, bu, po, per_file, art = summaries(m, d)
    season = "-".join(sorted(o["Stagione"].unique()))
    OUT_DIR.mkdir(exist_ok=True)
    xl = OUT_DIR / f"controllo_consegne_{season}.xlsx"
    write_excel(xl, tot, bu, po, per_file, art, m)
    html = OUT_DIR / f"controllo_consegne_{season}.html"
    pdf = OUT_DIR / f"controllo_consegne_{season}.pdf"
    write_pdf(html, pdf, season, tot, bu, po, per_file, art, m)

    pd.set_option("display.width", 200)
    print(tot.to_string(index=False))
    print(bu.to_string(index=False))
    print(po.to_string(index=False))
    print("Stati:", m["Stato"].value_counts().to_dict())
    print("Stagioni consegne:", d["Stagione"].value_counts().to_dict())
    print(xl, pdf, sep="\n")


if __name__ == "__main__":
    main()
